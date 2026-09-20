import pytest
import io
import docx
from engine.core.docx_report_generator import DocxReportGenerator


def test_docx_report_generation():
    """Validates generation of valid DOCX binary output and content structure."""
    cad_mock = {
        "filename": "Avionics_Chassis_v4.step",
        "mass_kg": 2.45,
        "volume_mm3": 907407.4,
        "bounding_box": {"lx": 180.0, "ly": 140.0, "lz": 75.0},
        "center_of_gravity": {"x": 90.0, "y": 70.0, "z": 32.5},
        "material": "Aluminium 6061-T6"
    }
    mission_mock = {
        "standard_id": "MIL-STD-810H",
        "category": "Metot 514.8 - Kategori 24",
        "duration_hours_per_axis": 1.0,
        "overall_g_rms": 7.67,
        "psd_breakpoints": [
            {"frequency_hz": 20.0, "psd_g2_hz": 0.04, "slope_db_octave": 0.0},
            {"frequency_hz": 1000.0, "psd_g2_hz": 0.04, "slope_db_octave": -6.0},
            {"frequency_hz": 2000.0, "psd_g2_hz": 0.01, "slope_db_octave": 0.0}
        ]
    }
    thermal_mock = {
        "temperature_profile": {"operational_high_c": 71.0, "operational_low_c": -40.0},
        "joint_thermal_analysis": {
            "delta_cte_ppm_per_k": 11.5,
            "ms_yield_hot": 0.65,
            "ms_separation_cold": 0.85,
            "preload_retention_percent": 88.5
        },
        "qualification_status": "PASS",
        "engineering_summary": "Termal genleşme ve ön yük kriterlerini sağlar."
    }
    shock_mock = {
        "pulse_shape": "Terminal Peak Sawtooth (TPS)",
        "peak_acceleration_g": 40.0,
        "duration_ms": 11.0,
        "f1_part_hz": 240.0,
        "f1_peak_response_g": 62.0,
        "total_shock_inertial_force_n": 1490.0,
        "force_per_bolt_n": 372.5,
        "fastener_safety_margins": {
            "bolt_tensile_margin_of_safety": 1.45,
            "bolt_shear_margin_of_safety": 2.10
        }
    }
    comp_mock = {
        "material_name": "AS4/3501-6 Carbon/Epoxy",
        "layup_sequence": "[0/45/-45/90]s",
        "num_plies": 8,
        "total_thickness_mm": 1.0,
        "tsai_wu_margin_of_safety": 0.42,
        "max_stress_margin_of_safety": 0.38,
        "critical_ply_index": 4
    }

    docx_bytes = DocxReportGenerator.generate_etp_docx(
        cad_data=cad_mock,
        mission_profile=mission_mock,
        thermal_data=thermal_mock,
        shock_data=shock_mock,
        composite_data=comp_mock,
        classification="HİZMETE ÖZEL / RESTRICTED"
    )

    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 5000

    # Verify document can be opened and parsed by docx library
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = " ".join([p.text for p in doc.paragraphs])
    assert "NUPER CITADEL" in paragraphs_text
    assert "MIL-STD-810H" in paragraphs_text
    assert len(doc.tables) >= 5
